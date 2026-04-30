"""Read an attachment's bytes and prepare them for the summarizer LLM.

Each file type has its own reader. The output (``PreparedContent``) is
content-type-agnostic from the summarizer's point of view: a text excerpt
and/or a list of images, plus light metadata (page/row counts).

Optional dependencies (graceful fallback):
  - pdfplumber   for PDF text extraction
  - python-docx  for .docx text extraction
  - openpyxl     for .xlsx reads
  - csv          (stdlib)

If a lib is missing we return a PreparedContent with an entry in ``notes``
explaining the gap, so the summarizer can still produce a metadata-only
summary instead of crashing the upload.
"""

from __future__ import annotations

import csv
import io
import logging
import mimetypes
import os
from dataclasses import dataclass, field
from typing import List, Optional, Tuple

from integrations.llm.provider import LLMImage

try:
    from logger import logger as _app_logger

    logger = _app_logger.getChild("partner.attachment.extract")
except Exception:
    logger = logging.getLogger(__name__)


# --- limits --------------------------------------------------------------
# Cap text passed to the summarizer so cost stays predictable. ~32k chars
# roughly equals ~8k tokens; cheap on gpt-4o-mini and well under context
# limits even on cheaper providers.
MAX_TEXT_CHARS = 32_000

# How many spreadsheet rows we sample for the summarizer. Header is always
# included; data rows beyond this are summarized as "(... N more rows)".
MAX_SPREADSHEET_ROWS = 50


# --- output --------------------------------------------------------------


@dataclass
class PreparedContent:
    classification: str  # 'pdf_report' | 'spreadsheet' | 'image' | 'doc' | 'text' | 'other'
    text_excerpt: Optional[str] = None
    images: List[LLMImage] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)
    page_count: Optional[int] = None
    row_count: Optional[int] = None
    col_count: Optional[int] = None
    truncated: bool = False


# --- type sniffing -------------------------------------------------------


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower().lstrip(".")


def _classify(mime: Optional[str], filename: str) -> str:
    m = (mime or "").lower().strip()
    e = _ext(filename)
    if m.startswith("image/") or e in {"jpg", "jpeg", "png", "gif", "webp", "bmp"}:
        return "image"
    if m == "application/pdf" or e == "pdf":
        return "pdf_report"
    if e in {"xlsx", "xlsm", "xls"} or "spreadsheetml" in m or m == "application/vnd.ms-excel":
        return "spreadsheet"
    if e == "csv" or m == "text/csv":
        return "spreadsheet"
    if e == "docx" or "wordprocessingml" in m:
        return "doc"
    if e in {"txt", "md", "markdown"} or m.startswith("text/"):
        return "text"
    return "other"


# --- readers -------------------------------------------------------------


def _read_text(blob: bytes, filename: str) -> Tuple[Optional[str], List[str], bool]:
    notes: List[str] = []
    truncated = False
    try:
        text = blob.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = blob.decode("latin-1")
            notes.append("decoded as latin-1 (utf-8 decode failed)")
        except Exception as e:
            notes.append(f"failed to decode text: {e}")
            return None, notes, truncated

    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        truncated = True
        notes.append(f"text truncated to {MAX_TEXT_CHARS} chars")
    return text, notes, truncated


def _read_csv(blob: bytes, filename: str) -> Tuple[Optional[str], List[str], int, int, bool]:
    notes: List[str] = []
    truncated = False
    try:
        text = blob.decode("utf-8")
    except UnicodeDecodeError:
        text = blob.decode("latin-1", errors="replace")
        notes.append("decoded as latin-1")

    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return None, notes, 0, 0, truncated

    header = rows[0]
    body = rows[1:]
    row_count = len(body)
    col_count = len(header)

    sample = body[:MAX_SPREADSHEET_ROWS]
    if len(body) > MAX_SPREADSHEET_ROWS:
        truncated = True
        notes.append(f"sampled first {MAX_SPREADSHEET_ROWS} of {row_count} rows")

    out_lines = ["| " + " | ".join(str(c) for c in header) + " |"]
    out_lines.append("| " + " | ".join("---" for _ in header) + " |")
    for r in sample:
        out_lines.append("| " + " | ".join(str(c) for c in r) + " |")

    excerpt = "\n".join(out_lines)
    if len(excerpt) > MAX_TEXT_CHARS:
        excerpt = excerpt[:MAX_TEXT_CHARS]
        truncated = True
        notes.append(f"excerpt truncated to {MAX_TEXT_CHARS} chars")
    return excerpt, notes, row_count, col_count, truncated


def _read_xlsx(blob: bytes, filename: str) -> Tuple[Optional[str], List[str], int, int, bool]:
    notes: List[str] = []
    truncated = False
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as e:
        notes.append(f"openpyxl not available; xlsx not read: {e}")
        return None, notes, 0, 0, truncated

    try:
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    except Exception as e:
        notes.append(f"failed to open workbook: {e}")
        return None, notes, 0, 0, truncated

    out_chunks: List[str] = []
    total_rows = 0
    max_cols = 0
    for ws in wb.worksheets:
        out_chunks.append(f"### Sheet: {ws.title}")
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        header = ["" if v is None else str(v) for v in header]
        max_cols = max(max_cols, len(header))
        out_chunks.append("| " + " | ".join(header) + " |")
        out_chunks.append("| " + " | ".join("---" for _ in header) + " |")
        sampled = 0
        sheet_rows = 0
        for row in rows_iter:
            sheet_rows += 1
            if sampled < MAX_SPREADSHEET_ROWS:
                cells = ["" if v is None else str(v) for v in row]
                out_chunks.append("| " + " | ".join(cells) + " |")
                sampled += 1
        total_rows += sheet_rows
        if sheet_rows > MAX_SPREADSHEET_ROWS:
            truncated = True
            out_chunks.append(f"(... {sheet_rows - MAX_SPREADSHEET_ROWS} more rows in '{ws.title}')")

    excerpt = "\n".join(out_chunks)
    if len(excerpt) > MAX_TEXT_CHARS:
        excerpt = excerpt[:MAX_TEXT_CHARS]
        truncated = True
        notes.append(f"excerpt truncated to {MAX_TEXT_CHARS} chars")
    return excerpt or None, notes, total_rows, max_cols, truncated


def _read_docx(blob: bytes, filename: str) -> Tuple[Optional[str], List[str], bool]:
    notes: List[str] = []
    truncated = False
    try:
        import docx  # type: ignore  # python-docx
    except Exception as e:
        notes.append(f"python-docx not available; docx not read: {e}")
        return None, notes, truncated

    try:
        document = docx.Document(io.BytesIO(blob))
    except Exception as e:
        notes.append(f"failed to open docx: {e}")
        return None, notes, truncated

    paragraphs = [p.text for p in document.paragraphs if p.text]
    text = "\n".join(paragraphs)

    # Append simple table dumps so we don't drop tabular content.
    for ti, table in enumerate(document.tables, 1):
        text += f"\n\n[table {ti}]\n"
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            text += " | ".join(cells) + "\n"

    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        truncated = True
        notes.append(f"text truncated to {MAX_TEXT_CHARS} chars")
    return text or None, notes, truncated


def _read_pdf(blob: bytes, filename: str) -> Tuple[Optional[str], List[str], int, bool]:
    notes: List[str] = []
    truncated = False
    page_count = 0
    try:
        import pdfplumber  # type: ignore
    except Exception as e:
        notes.append(f"pdfplumber not available; pdf not read: {e}")
        return None, notes, 0, truncated

    try:
        with pdfplumber.open(io.BytesIO(blob)) as pdf:
            page_count = len(pdf.pages)
            chunks: List[str] = []
            for idx, page in enumerate(pdf.pages, 1):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                if text:
                    chunks.append(f"--- page {idx} ---\n{text}")
                if sum(len(c) for c in chunks) > MAX_TEXT_CHARS:
                    truncated = True
                    notes.append(
                        f"stopped reading at page {idx}; excerpt cap reached "
                        f"({MAX_TEXT_CHARS} chars)"
                    )
                    break
            joined = "\n\n".join(chunks)
    except Exception as e:
        notes.append(f"failed to read pdf: {e}")
        return None, notes, page_count, truncated

    if not joined:
        notes.append(
            "no extractable text -- pdf may be scanned/image-based; OCR not enabled in v1"
        )
        return None, notes, page_count, truncated

    if len(joined) > MAX_TEXT_CHARS:
        joined = joined[:MAX_TEXT_CHARS]
        truncated = True
    return joined, notes, page_count, truncated


# --- main entry point ----------------------------------------------------


def prepare_content(
    blob: bytes,
    *,
    filename: str,
    content_type: Optional[str] = None,
) -> PreparedContent:
    """Convert raw bytes into something the summarizer can consume."""
    classification = _classify(content_type, filename)

    # Image: pass bytes straight to the vision model.
    if classification == "image":
        mime = (content_type or "").strip()
        if not mime:
            guessed, _ = mimetypes.guess_type(filename)
            mime = guessed or "image/jpeg"
        return PreparedContent(
            classification="image",
            images=[LLMImage(bytes_=blob, mime=mime, filename=filename)],
        )

    if classification == "pdf_report":
        text, notes, pages, trunc = _read_pdf(blob, filename)
        return PreparedContent(
            classification="pdf_report",
            text_excerpt=text,
            notes=notes,
            page_count=pages or None,
            truncated=trunc,
        )

    if classification == "spreadsheet":
        e = _ext(filename)
        if e == "csv":
            text, notes, rows, cols, trunc = _read_csv(blob, filename)
        else:
            text, notes, rows, cols, trunc = _read_xlsx(blob, filename)
        return PreparedContent(
            classification="spreadsheet",
            text_excerpt=text,
            notes=notes,
            row_count=rows or None,
            col_count=cols or None,
            truncated=trunc,
        )

    if classification == "doc":
        text, notes, trunc = _read_docx(blob, filename)
        return PreparedContent(
            classification="doc",
            text_excerpt=text,
            notes=notes,
            truncated=trunc,
        )

    if classification == "text":
        text, notes, trunc = _read_text(blob, filename)
        return PreparedContent(
            classification="text",
            text_excerpt=text,
            notes=notes,
            truncated=trunc,
        )

    return PreparedContent(
        classification="other",
        notes=[f"unsupported content type '{content_type}' / extension '{_ext(filename)}'"],
    )


__all__ = ["PreparedContent", "prepare_content"]
